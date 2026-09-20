from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment

DESKTOP = Path(r"C:\Users\lironghua\Desktop")
SOURCE = DESKTOP / "publication data.xls"
OUTPUT = DESKTOP / "sample_comparison.xlsx"
COUNTRIES = ["Brazil", "India"]
SUBTYPES = ["Flash flood", "Flood (General)", "Riverine flood"]


def read_data(source):
    events = pd.read_excel(source, sheet_name="EMdat", engine="xlrd")
    gdp = pd.read_excel(source, sheet_name="GDP", engine="xlrd")
    events.columns = events.columns.astype(str).str.strip()
    gdp.columns = gdp.columns.astype(str).str.strip()
    damage = "Total Damage ('000 US$)"
    for col in ["Country", "Disaster Type", "Disaster Subtype"]:
        events[col] = events[col].astype(str).str.strip()
    for col in [damage, "Start Year"]:
        events[col] = pd.to_numeric(events[col], errors="coerce")
    selected = (
        events["Country"].isin(COUNTRIES)
        & events["Disaster Type"].eq("Flood")
        & events["Disaster Subtype"].isin(SUBTYPES)
        & events["Start Year"].between(2000, 2024)
        & events[damage].gt(0)
    )
    data = events.loc[selected].copy()
    if data.groupby("Country").size().to_dict() != {"Brazil": 51, "India": 68}:
        raise ValueError("Expected 51 Brazil events and 68 India events. Check the sample.")
    if data["DisNo."].duplicated().any():
        raise ValueError("Duplicate event IDs. Check the source data.")
    data["Excel row"] = data.index + 2
    data["GDP (USD)"] = np.nan
    for country in COUNTRIES:
        row = gdp.loc[gdp["Country Name"].astype(str).str.strip().eq(country)]
        if len(row) != 1 or row.iloc[0]["Indicator Name"] != "GDP (current US$)":
            raise ValueError(f"Check the GDP row and units for {country}.")
        annual = {int(y): float(row.iloc[0][y]) for y in gdp.columns if y.isdigit()}
        mask = data["Country"].eq(country)
        data.loc[mask, "GDP (USD)"] = data.loc[mask, "Start Year"].map(annual)
    if not (np.isfinite(data["GDP (USD)"]) & data["GDP (USD)"].gt(0)).all():
        raise ValueError("Missing or invalid GDP values.")

    data["Damage (USD)"] = data[damage] * 1000
    data["Loss/GDP"] = data["Damage (USD)"] / data["GDP (USD)"]
    if not np.isfinite(data["Loss/GDP"]).all():
        raise ValueError("Invalid loss ratios.")
    columns = ["Country", "DisNo.", "Start Year", "Disaster Subtype", "Excel row",
               damage, "GDP (USD)", "Damage (USD)", "Loss/GDP"]
    return data[columns].sort_values(["Country", "Start Year", "DisNo."])


def make_summary(data):
    results = {}
    for country in COUNTRIES:
        sample = data.loc[data["Country"].eq(country)]
        x = sample["Loss/GDP"]
        results[country] = {
            "Study period": "2000-2024",
            "Years with positive-loss records":
                f"{int(sample['Start Year'].min())}-{int(sample['Start Year'].max())}",
            "Flood subtypes": "; ".join(SUBTYPES),
            "Positive-loss events (n)": len(sample),
            "Minimum loss/GDP": x.min(),
            "P25 loss/GDP": x.quantile(0.25, interpolation="linear"),
            "Median loss/GDP": x.median(),
            "P75 loss/GDP": x.quantile(0.75, interpolation="linear"),
            "P90 loss/GDP": x.quantile(0.90, interpolation="linear"),
            "P95 loss/GDP": x.quantile(0.95, interpolation="linear"),
            "Maximum loss/GDP": x.max()
        }
    return pd.DataFrame(results).rename_axis("Measure").reset_index()


def main(source=SOURCE, output=OUTPUT):
    data = read_data(source)
    summary = make_summary(data)
    notes = [
        "Source: publication data.xls, EMdat and GDP sheets.",
        "Sample: Flood events in 2000-2024, using the three listed subtypes and positive reported losses. Coastal flood is excluded.",
        "Missing and zero losses are excluded. Missing losses are not treated as zero.",
        "Loss/GDP = Total Damage ('000 US$) * 1000 / GDP (current US$), matched by country and event start year.",
        "Loss/GDP values are displayed as percentages. Quantiles use linear interpolation (Excel PERCENTILE.INC).",
        "Year ranges show the earliest and latest positive-loss records. They do not imply a record in every year or a shorter study period.",
        "Statistics describe individual recorded events, not annual aggregate losses. Rerun Python to refresh the results."
    ]
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        data.to_excel(writer, sheet_name="Events", index=False)
        sheet = writer.sheets["Summary"]
        sheet.column_dimensions["A"].width = 39
        for col in ["B", "C"]:
            sheet.column_dimensions[col].width = 32
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            sheet.row_dimensions[row[0].row].height = 23
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[4].height = 48
        for row in sheet.iter_rows(min_row=6, max_row=sheet.max_row, min_col=2):
            for cell in row:
                cell.number_format = "0.000000%"
        for i, note in enumerate(notes, start=len(summary) + 4):
            sheet.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
            sheet.cell(i, 1, note).alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[i].height = 32
        sheet = writer.sheets["Events"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 32
        for column in sheet.columns:
            header = column[0].value
            sheet.column_dimensions[column[0].column_letter].width = 25 if "USD" in header or "US$" in header else 19
            column[0].alignment = Alignment(wrap_text=True, vertical="center")
            for cell in column[1:]:
                if header == "Loss/GDP":
                    cell.number_format = "0.000000%"
                elif "USD" in header or "US$" in header:
                    cell.number_format = "#,##0.00"
    print(summary.to_string(index=False))
    print(f"Saved: {output}")


if __name__ == "__main__":
    main()
